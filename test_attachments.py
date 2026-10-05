"""Attachment extraction and chat-context regression checks."""

import json
import base64
import queue
import runpy
import sqlite3
import sys
import tempfile
import threading
import unittest
from contextlib import contextmanager
from pathlib import Path
from unittest.mock import patch

from attachments import (copy_sources, process_attachment, relevant_chunks,
                         validate_sources)


CHAT_DIR = Path(__file__).resolve().parent
sys.path.insert(0, str(CHAT_DIR))
MODULE = runpy.run_path(str(CHAT_DIR / "local_chat.pyw"), run_name="test_import")
LocalChat = MODULE["LocalChat"]


class AttachmentTests(unittest.TestCase):
    @contextmanager
    def attachment_database(self, root, copied, store=None):
        database = root / "history.sqlite3"
        store = root / "attachments" if store is None else store

        @contextmanager
        def managed_connection():
            db = sqlite3.connect(database)
            db.execute("PRAGMA foreign_keys = ON")
            try:
                with db:
                    yield db
            finally:
                db.close()

        globals_patch = {"db_connect": managed_connection,
                         "ATTACHMENTS_DIR": store}
        with patch.dict(MODULE["init_db"].__globals__, globals_patch):
            MODULE["init_db"]()
            with managed_connection() as db:
                db.execute("INSERT INTO chats(id,title,updated_at) VALUES (1,'Test','now')")
                db.execute("INSERT INTO messages(id,chat_id,role,content) VALUES (1,1,'user','Прочитай файл')")
                db.execute("INSERT INTO messages(id,chat_id,role,content) VALUES (2,1,'user','Какое число?')")
                db.execute(
                    "INSERT INTO attachments(message_id,name,kind,path) VALUES (1,?,?,?)",
                    (copied["name"], copied["kind"], copied["path"]),
                )
            app = LocalChat.__new__(LocalChat)
            app.events = queue.Queue()
            app.usage_report = None
            app.chat_id = 1
            yield app, managed_connection

    def test_image_and_document_formats(self):
        from PIL import Image
        import pymupdf
        from docx import Document
        from openpyxl import Workbook
        from pptx import Presentation

        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            image = root / "photo.png"
            Image.new("RGB", (2400, 1200), "red").save(image)
            photo = process_attachment(image, "image")
            with Image.open(photo["images"][0]) as prepared:
                self.assertLessEqual(max(prepared.size), 1600)

            pdf_path = root / "scan.pdf"
            pdf = pymupdf.open()
            pdf.new_page(width=400, height=300)
            pdf.save(pdf_path)
            pdf.close()
            pdf_result = process_attachment(pdf_path, "document")
            self.assertEqual(pdf_result["scanned_pages"], [1])
            self.assertTrue(pdf_result["images"][0].is_file())

            docx_path = root / "notes.docx"
            doc = Document()
            doc.add_paragraph("Важный вывод")
            table = doc.add_table(rows=1, cols=2)
            table.cell(0, 0).text = "A"
            table.cell(0, 1).text = "42"
            doc.save(docx_path)
            doc_text = process_attachment(docx_path, "document")["text"]
            self.assertIn("Важный вывод", doc_text)
            self.assertIn("A | 42", doc_text)

            xlsx_path = root / "data.xlsx"
            workbook = Workbook()
            workbook.active["A1"] = "Итог"
            workbook.active["B1"] = 42
            workbook.save(xlsx_path)
            sheet_text = process_attachment(xlsx_path, "document")["text"]
            self.assertIn("B1=42", sheet_text)

            pptx_path = root / "slides.pptx"
            presentation = Presentation()
            slide = presentation.slides.add_slide(presentation.slide_layouts[6])
            box = slide.shapes.add_textbox(0, 0, 3000000, 500000)
            box.text = "Слайд с выводом"
            presentation.save(pptx_path)
            self.assertIn("Слайд с выводом",
                          process_attachment(pptx_path, "document")["text"])

    def test_copied_attachment_is_used_for_followup(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "source.txt"
            source.write_text("Секретное число: 42", encoding="utf-8")
            store = root / "attachments"
            copied = copy_sources(validate_sources([source]), store, 1)[0]
            source.unlink()
            with self.attachment_database(root, copied) as (app, _):
                messages = [
                    {"role": "user", "content": "Прочитай файл", "message_id": 1},
                    {"role": "user", "content": "Какое число?", "message_id": 2},
                ]
                prepared = app.prepare_chat_messages(
                    messages, 8192, 2048, 3, threading.Event())
                self.assertIn("Секретное число: 42", prepared[1]["content"])
                self.assertNotIn("message_id", prepared[1])
                self.assertEqual(prepared[2]["content"], "Какое число?")
                older = app.prepare_chat_messages(
                    [{"role": "user", "content": "source.txt: какое число?",
                      "message_id": 2}], 8192, 2048, 4, threading.Event())
                self.assertIn("Секретное число: 42", older[-1]["content"])

    def test_attachment_metadata_handles_an_aliased_storage_path(self):
        from PIL import Image

        for kind in ("document", "image"):
            with self.subTest(kind=kind), tempfile.TemporaryDirectory() as directory:
                root = Path(directory)
                store = root / "attachments"
                if kind == "document":
                    source = root / "notes.txt"
                    source.write_text("Значение: 42", encoding="utf-8")
                else:
                    source = root / "photo.png"
                    Image.new("RGB", (40, 30), "blue").save(source)
                copied = copy_sources(validate_sources([source]), store, 1)[0]
                alias = store / ".." / "attachments"
                with self.attachment_database(root, copied, store=alias) as (app, connect):
                    prepared = app.prepare_chat_messages(
                        [{"role": "user", "content": "Прочитай файл", "message_id": 1}],
                        8192, 2048, 3, threading.Event())
                    if kind == "document":
                        self.assertIn("Значение: 42", prepared[-1]["content"])
                    else:
                        self.assertTrue(base64.b64decode(prepared[-1]["images"][0]).startswith(b"\xff\xd8"))
                    with connect() as db:
                        metadata = json.loads(db.execute(
                            "SELECT metadata FROM attachments").fetchone()[0])
                    paths = metadata["images"] + ([metadata["text_path"]] if metadata["text_path"] else [])
                    self.assertTrue(paths)
                    for relative in paths:
                        self.assertFalse(Path(relative).is_absolute())
                        self.assertNotIn("..", Path(relative).parts)
                        self.assertTrue((store / relative).is_file())

    def test_photo_payload_contains_jpeg_and_scan_is_cached(self):
        from PIL import Image
        import pymupdf

        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            store = root / "attachments"
            photo = root / "photo.png"
            Image.new("RGB", (120, 80), "blue").save(photo)
            copied_photo = copy_sources(validate_sources([photo]), store, 1)[0]
            with self.attachment_database(root, copied_photo) as (app, _):
                prepared = app.prepare_chat_messages(
                    [{"role": "user", "content": "Что на фото?", "message_id": 1}],
                    8192, 2048, 3, threading.Event())
                self.assertTrue(base64.b64decode(prepared[-1]["images"][0]).startswith(b"\xff\xd8"))

        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            store = root / "attachments"
            pdf_path = root / "scan.pdf"
            pdf = pymupdf.open()
            pdf.new_page(width=400, height=300)
            pdf.save(pdf_path)
            pdf.close()
            copied_pdf = copy_sources(validate_sources([pdf_path]), store, 1)[0]
            with self.attachment_database(root, copied_pdf) as (app, _):
                calls = []

                def fake_ocr(payload, *_args, **_kwargs):
                    calls.append(payload)
                    return "Номер счёта 42", "stop"

                app.analysis_request_with_retries = fake_ocr
                messages = [{"role": "user", "content": "Какой номер счёта?", "message_id": 1}]
                first = app.prepare_chat_messages(messages, 8192, 2048, 3, threading.Event())
                second = app.prepare_chat_messages(messages, 8192, 2048, 4, threading.Event())
                self.assertEqual(len(calls), 1)
                self.assertIn("Номер счёта 42", first[-1]["content"])
                self.assertIn("Номер счёта 42", second[-1]["content"])

    def test_long_document_marks_selected_fragments(self):
        text = "A" * 2000 + " искомый раздел " + "B" * 2000
        result = relevant_chunks(text, "искомый раздел", 1900)
        self.assertIn("искомый раздел", result)
        self.assertIn("Показано", result)

    def test_clear_all_data_removes_copies_but_keeps_original(self):
        cleanup = runpy.run_path(str(CHAT_DIR / "clear_history.pyw"),
                                 run_name="test_import")["purge_local_data"]
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            app_dir = root / "Chat"
            app_dir.mkdir()
            original = root / "original.txt"
            original.write_text("Исходник", encoding="utf-8")
            copy = app_dir / "attachments" / "1" / "copy.txt"
            copy.parent.mkdir(parents=True)
            copy.write_text("Копия", encoding="utf-8")
            db = sqlite3.connect(app_dir / "history.sqlite3")
            db.execute("CREATE TABLE chats(id INTEGER PRIMARY KEY)")
            db.execute("CREATE TABLE messages(id INTEGER PRIMARY KEY)")
            db.execute("CREATE TABLE settings(key TEXT PRIMARY KEY)")
            db.commit()
            db.close()
            cleanup(app_dir)
            self.assertTrue(original.is_file())
            self.assertFalse((app_dir / "attachments").exists())
            self.assertFalse((app_dir / "history.sqlite3").exists())

    def test_attachment_controls_build_in_tk_window(self):
        import tkinter as tk

        with tempfile.TemporaryDirectory() as directory:
            database = Path(directory) / "history.sqlite3"
            store = Path(directory) / "attachments"
            source = Path(directory) / "note.txt"
            source.write_text("Проверка вложения", encoding="utf-8")

            @contextmanager
            def managed_connection():
                db = sqlite3.connect(database)
                db.execute("PRAGMA foreign_keys = ON")
                try:
                    with db:
                        yield db
                finally:
                    db.close()

            with patch.dict(MODULE["init_db"].__globals__,
                            {"db_connect": managed_connection,
                             "ATTACHMENTS_DIR": store}), patch.object(
                                LocalChat, "ensure_server", return_value=None):
                MODULE["init_db"]()
                root = tk.Tk()
                root.withdraw()
                try:
                    app = LocalChat(root)
                    self.assertEqual(app.attach_button.cget("text"), "Прикрепить…")
                    self.assertEqual(app.pending_attachments, [])
                    captured = []

                    def fake_generate(_self, messages, *_args):
                        captured.extend(messages)

                    app.ready = True
                    app.pending_attachments = [str(source)]
                    app.refresh_attachment_list()
                    with patch.object(LocalChat, "generate", fake_generate):
                        app.send()
                        app.generation_thread.join(timeout=2)
                    self.assertTrue(captured)
                    self.assertEqual(app.pending_attachments, [])
                    with managed_connection() as db:
                        rows = db.execute("SELECT name, path FROM attachments").fetchall()
                    self.assertEqual(rows[0][0], "note.txt")
                    self.assertTrue((store / rows[0][1]).is_file())
                finally:
                    root.destroy()


if __name__ == "__main__":
    unittest.main()
